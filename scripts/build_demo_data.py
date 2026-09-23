#!/usr/bin/env python3
"""Build a privacy-reduced demo catalog from supplier ZIP reports.

The output keeps monthly SKU aggregates only. It drops invoice numbers,
document descriptions, warehouse rows, and any customer-identifying fields.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import statistics
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


CYR = dict(zip(
    "абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
    ["a", "b", "v", "g", "d", "e", "yo", "zh", "z", "i", "y", "k", "l", "m", "n", "o", "p", "r", "s", "t", "u", "f", "h", "ts", "ch", "sh", "shch", "", "y", "", "e", "yu", "ya"],
))
CYR.update(dict(zip("АБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯ", [CYR[c.lower()] for c in "АБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯ"])))
CYR.update({"Қ": "Q", "қ": "q", "Ғ": "Gh", "ғ": "gh", "Ң": "Ng", "ң": "ng", "Ұ": "U", "ұ": "u", "Ү": "U", "ү": "u", "Һ": "H", "һ": "h", "Ө": "O", "ө": "o", "Ә": "A", "ә": "a", "І": "I", "і": "i"})

MONTHS = {
    "yanv": 1, "janv": 1, "fevr": 2, "fev": 2, "mart": 3, "mar": 3,
    "aprel": 4, "apr": 4, "may": 5, "iyun": 6, "iyul": 7,
    "iyul": 7, "avgust": 8, "avg": 8, "sent": 9, "sentyabr": 9,
    "sen": 9, "okt": 10, "oktyabr": 10, "noyab": 11, "noya": 11,
    "dek": 12,
}


def latin(value: object) -> str:
    return "".join(CYR.get(char, char) for char in str(value))


def normalized(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", latin(value).lower()).strip()


def code(value: object) -> str:
    return str(value or "").strip().rstrip("_").strip()


def number(value: object, default: float = 0.0) -> float:
    if value is None or isinstance(value, bool):
        return default
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        text = str(value).strip().replace(" ", "").replace(",", ".")
        try:
            result = float(text)
        except ValueError:
            return default
    return result if math.isfinite(result) else default


def parse_month(value: object) -> str | None:
    text = normalized(value)
    match = re.search(r"([a-z]+)\s+(20\d{2})", text)
    if not match:
        return None
    month_name, year = match.groups()
    month = next((num for prefix, num in MONTHS.items() if month_name.startswith(prefix)), None)
    if not month:
        return None
    return f"{year}-{month:02d}"


def open_first_sheet(payload: bytes):
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    return workbook, workbook.worksheets[0]


def source_files(archive: Path) -> dict[str, tuple[str, bytes]]:
    found: dict[str, tuple[str, bytes]] = {}
    with ZipFile(archive) as zipped:
        for member in zipped.namelist():
            if not member.lower().endswith(".xlsx"):
                continue
            key = normalized(Path(member).name)
            kind = None
            if "moq" in key:
                kind = "moq"
            elif "dinamika prodazh" in key:
                kind = "transactions"
            elif "ezhemesyachnye prodazhi" in key and "kol" in key:
                kind = "sales"
            elif "ezhemesyachnye ostatki" in key:
                kind = "monthly_stock"
            elif "tovar v puti" in key:
                kind = "system_summary"
            elif "put iek" in key:
                kind = "iek_inbound"
            if kind and kind not in found:
                found[kind] = (Path(member).name, zipped.read(member))
    return found


def header_row(sheet, required_terms: tuple[str, ...], limit: int = 8):
    for row_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(limit, sheet.max_row), values_only=True), 1):
        labels = [normalized(value) for value in row]
        joined = " | ".join(labels)
        if all(term in joined for term in required_terms):
            return row_index, [str(value or "") for value in row]
    raise ValueError(f"Could not find worksheet header containing: {required_terms}")


def column_index(headers: list[str], *terms: str) -> int | None:
    for index, header in enumerate(headers):
        label = normalized(header)
        if all(term in label for term in terms):
            return index
    return None


def product_base(vendor: str, sku: str, name: str) -> dict:
    return {
        "sku": sku,
        "supplier": vendor,
        "name": name.strip(),
        "supplier_sku": "",
        "category": "Без категории",
        "category_known": False,
        "unit": "шт.",
        "moq": 1.0,
        "moq_known": False,
        "current_stock": 0.0,
        "stock_known": False,
        "stock_snapshot": "",
        "stock_basis": "Остаток не найден в выгрузке",
        "in_transit": 0.0,
        "inbound_known": False,
        "monthly_sales": {},
        "one_off_adjustments": {},
        "detected_large_orders": 0,
    }


def extract_moq(files, products):
    name, payload = files["moq"]
    workbook, sheet = open_first_sheet(payload)
    try:
        header_line, headers = header_row(sheet, ("kod",), limit=4)
        code_col = column_index(headers, "kod")
        if code_col is None:
            code_col = column_index(headers, "nomenklatura")
        moq_col = next((i for i, h in enumerate(headers) if any(word in normalized(h) for word in ("kratnost", "razr"))), None)
        article_col = next((i for i, h in enumerate(headers) if "artikul" in normalized(h)), None)
        name_col = next((i for i, h in enumerate(headers) if any(word in normalized(h) for word in ("nomenklatura", "naimenovanie")) and i != code_col), None)
        if moq_col is None or code_col is None:
            raise ValueError(f"Missing SKU/MOQ columns in {name}")
        for row in sheet.iter_rows(min_row=header_line + 1, values_only=True):
            sku = code(row[code_col] if code_col < len(row) else None)
            if not sku:
                continue
            product = products.setdefault(sku, product_base(files["vendor"], sku, str(row[name_col] or "") if name_col is not None else ""))
            raw_moq = row[moq_col] if moq_col < len(row) else None
            product["moq"] = max(1.0, number(raw_moq, 1.0))
            product["moq_known"] = raw_moq is not None and number(raw_moq, 0.0) > 0
            if article_col is not None and article_col < len(row):
                product["supplier_sku"] = str(row[article_col] or "").strip()
    finally:
        workbook.close()


def extract_monthly_sales(files, products):
    name, payload = files["sales"]
    workbook, sheet = open_first_sheet(payload)
    try:
        header_line, headers = header_row(sheet, ("nomenklatura", "kod"), limit=4)
        code_col = next((i for i, h in enumerate(headers) if "kod" in normalized(h)), None)
        name_col = next((i for i, h in enumerate(headers) if "nomenklatura" in normalized(h)), None)
        article_col = next((i for i, h in enumerate(headers) if "artikul" in normalized(h)), None)
        month_cols = [(i, parse_month(header)) for i, header in enumerate(headers)]
        month_cols = [(i, month) for i, month in month_cols if month]
        if code_col is None or not month_cols:
            raise ValueError(f"Missing monthly sales columns in {name}")
        for row in sheet.iter_rows(min_row=header_line + 1, values_only=True):
            sku = code(row[code_col] if code_col < len(row) else None)
            if not sku:
                continue
            product_name = str(row[name_col] or "").strip() if name_col is not None and name_col < len(row) else ""
            product = products.setdefault(sku, product_base(files["vendor"], sku, product_name))
            if not product["name"]:
                product["name"] = product_name
            if article_col is not None and article_col < len(row):
                product["supplier_sku"] = str(row[article_col] or "").strip() or product["supplier_sku"]
            sales = product["monthly_sales"]
            for col, month in month_cols:
                qty = number(row[col] if col < len(row) else None)
                # Monthly summaries are net sales. Negative return/reversal months
                # cannot create a negative replenishment forecast.
                sales[month] = sales.get(month, 0.0) + max(0.0, qty)
    finally:
        workbook.close()


def extract_system_summary(files, products):
    if "system_summary" not in files:
        return
    name, payload = files["system_summary"]
    workbook, sheet = open_first_sheet(payload)
    try:
        header_line, headers = header_row(sheet, ("naimenovanie", "ostatok"), limit=8)
        code_col = column_index(headers, "kod", "1s")
        stock_col = next((i for i, h in enumerate(headers) if "svobodnyy ostatok" in normalized(h)), None)
        if stock_col is None:
            stock_col = next((i for i, h in enumerate(headers) if normalized(h).strip() == "ostatok"), None)
        inbound_col = next((i for i, h in enumerate(headers) if "puti" in normalized(h)), None)
        category_col = next((i for i, h in enumerate(headers) if "kategoriya" in normalized(h)), None)
        article_col = next((i for i, h in enumerate(headers) if "artikul postavshchika" in normalized(h)), None)
        name_col = column_index(headers, "naimenovanie")
        if code_col is None:
            raise ValueError(f"Missing SKU column in {name}")
        for row in sheet.iter_rows(min_row=header_line + 1, values_only=True):
            sku = code(row[code_col] if code_col < len(row) else None)
            if not sku:
                continue
            product_name = str(row[name_col] or "").strip() if name_col is not None and name_col < len(row) else ""
            product = products.setdefault(sku, product_base(files["vendor"], sku, product_name))
            product["name"] = product["name"] or product_name
            if stock_col is not None and stock_col < len(row):
                raw_stock = row[stock_col]
                if raw_stock is not None:
                    product["current_stock"] = max(0.0, number(raw_stock))
                    product["stock_known"] = True
                    product["stock_basis"] = "Свободный остаток из сводного отчёта поставщика"
                    product["stock_snapshot"] = "2026-09-22"
            if inbound_col is not None and inbound_col < len(row):
                raw_inbound = row[inbound_col]
                # The supplier summary is a full SKU list; an empty cell means
                # no quantity is listed as in transit for that item.
                product["in_transit"] = max(0.0, number(raw_inbound))
                product["inbound_known"] = True
            if category_col is not None and category_col < len(row) and row[category_col] is not None:
                product["category"] = str(row[category_col]).strip()
                product["category_known"] = product["category"] not in ("", "0", "0.0")
            if article_col is not None and article_col < len(row):
                product["supplier_sku"] = str(row[article_col] or "").strip() or product["supplier_sku"]
    finally:
        workbook.close()


def extract_monthly_stock(files, products):
    if "monthly_stock" not in files:
        return
    name, payload = files["monthly_stock"]
    workbook, sheet = open_first_sheet(payload)
    try:
        header_line, headers = header_row(sheet, ("kod",), limit=4)
        code_col = next((i for i, h in enumerate(headers) if "kod" in normalized(h)), None)
        name_col = next((i for i, h in enumerate(headers) if "nomenklatura" in normalized(h)), None)
        unit_col = next((i for i, h in enumerate(headers) if normalized(h).strip() in ("ed", "ed izm")), None)
        month_cols = [(i, parse_month(header)) for i, header in enumerate(headers)]
        month_cols = [(i, month) for i, month in month_cols if month]
        if code_col is None or not month_cols:
            raise ValueError(f"Missing month-end inventory columns in {name}")
        latest_col, latest_month = max(month_cols, key=lambda pair: pair[1])
        for row in sheet.iter_rows(min_row=header_line + 1, values_only=True):
            sku = code(row[code_col] if code_col < len(row) else None)
            if not sku:
                continue
            product_name = str(row[name_col] or "").strip() if name_col is not None and name_col < len(row) else ""
            product = products.setdefault(sku, product_base(files["vendor"], sku, product_name))
            product["name"] = product["name"] or product_name
            if unit_col is not None and unit_col < len(row):
                product["unit"] = str(row[unit_col] or "шт.").strip()
            # System Electric's report has a newer free-stock snapshot and takes
            # precedence over the monthly balance history.
            if product["stock_basis"] == "Остаток не найден в выгрузке":
                product["current_stock"] = max(0.0, number(row[latest_col] if latest_col < len(row) else None))
                product["stock_known"] = True
                product["stock_snapshot"] = latest_month
                product["stock_basis"] = f"Остаток на начало {latest_month} (помесячная выгрузка)"
    finally:
        workbook.close()


def extract_iek_inbound(files, products):
    if "iek_inbound" not in files:
        return
    name, payload = files["iek_inbound"]
    workbook, sheet = open_first_sheet(payload)
    try:
        header_line, headers = header_row(sheet, ("kod", "artikul"), limit=4)
        code_col = next((i for i, h in enumerate(headers) if "kod" in normalized(h)), None)
        name_col = next((i for i, h in enumerate(headers) if "naimenovanie" in normalized(h)), None)
        article_col = next((i for i, h in enumerate(headers) if "artikul" in normalized(h)), None)
        inbound_cols = [i for i, header in enumerate(headers) if i not in (code_col, article_col, name_col)]
        if code_col is None:
            raise ValueError(f"Missing SKU column in {name}")
        for row in sheet.iter_rows(min_row=header_line + 1, values_only=True):
            sku = code(row[code_col] if code_col < len(row) else None)
            if not sku:
                continue
            product = products.setdefault(sku, product_base(files["vendor"], sku, str(row[name_col] or "") if name_col is not None and name_col < len(row) else ""))
            product["name"] = product["name"] or (str(row[name_col] or "").strip() if name_col is not None and name_col < len(row) else "")
            if article_col is not None and article_col < len(row):
                product["supplier_sku"] = str(row[article_col] or "").strip() or product["supplier_sku"]
            product["in_transit"] = sum(max(0.0, number(row[i] if i < len(row) else None)) for i in inbound_cols)
            product["inbound_known"] = True
    finally:
        workbook.close()


def parse_transaction_month(value: object) -> str | None:
    if isinstance(value, (date, datetime)):
        return f"{value.year:04d}-{value.month:02d}"
    text = str(value or "").strip()
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(text, fmt)
            return f"{parsed.year:04d}-{parsed.month:02d}"
        except ValueError:
            pass
    return None


def percentile(values: list[float], fraction: float) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    position = (len(ordered) - 1) * fraction
    lower = int(position)
    upper = min(lower + 1, len(ordered) - 1)
    weight = position - lower
    return ordered[lower] * (1 - weight) + ordered[upper] * weight


def detect_document_outliers(files, products):
    if "transactions" not in files:
        return
    name, payload = files["transactions"]
    workbook, sheet = open_first_sheet(payload)
    document_events: dict[tuple[str, str, str], float] = defaultdict(float)
    try:
        header_line, headers = header_row(sheet, ("data", "kod", "kolichestvo"), limit=4)
        date_col = next((i for i, h in enumerate(headers) if normalized(h) == "data"), None)
        document_col = next((i for i, h in enumerate(headers) if "nomer" in normalized(h)), None)
        code_col = next((i for i, h in enumerate(headers) if normalized(h) == "kod"), None)
        quantity_col = next((i for i, h in enumerate(headers) if "kolichestvo" in normalized(h)), None)
        description_col = next((i for i, h in enumerate(headers) if "dokument" in normalized(h)), None)
        if None in (date_col, code_col, quantity_col):
            return
        for row in sheet.iter_rows(min_row=header_line + 1, values_only=True):
            sku = code(row[code_col] if code_col < len(row) else None)
            month = parse_transaction_month(row[date_col] if date_col < len(row) else None)
            qty = number(row[quantity_col] if quantity_col < len(row) else None)
            if not sku or not month or qty >= 0:
                continue
            description = normalized(row[description_col]) if description_col is not None and description_col < len(row) else ""
            # The provided ledgers represent outbound sales as negative quantities.
            if description and not any(word in description for word in ("rashodnaya nakladnaya", "realizatsiya", "prodazha")):
                continue
            document = str(row[document_col] if document_col is not None and document_col < len(row) else "").strip()
            if not document:
                continue
            document_events[(sku, document, month)] += abs(qty)
    finally:
        workbook.close()

    by_sku: dict[str, list[tuple[str, float]]] = defaultdict(list)
    for (sku, _document, month), qty in document_events.items():
        by_sku[sku].append((month, qty))
    for sku, events in by_sku.items():
        if len(events) < 5 or sku not in products:
            continue
        event_sizes = [qty for _, qty in events]
        median_event = statistics.median(event_sizes)
        q1, q3 = percentile(event_sizes, 0.25), percentile(event_sizes, 0.75)
        monthly = [v for v in products[sku]["monthly_sales"].values() if v > 0]
        monthly_typical = statistics.median(monthly) if monthly else 0.0
        threshold = max(median_event * 5.0, q3 + 3.0 * (q3 - q1), monthly_typical * 2.0)
        flagged: dict[str, float] = defaultdict(float)
        for month, qty in events:
            if qty > threshold:
                flagged[month] += qty
        product = products[sku]
        for month, qty in flagged.items():
            raw = max(0.0, number(product["monthly_sales"].get(month)))
            removed = min(raw, qty)
            if removed > 0:
                product["one_off_adjustments"][month] = round(removed, 3)
        product["detected_large_orders"] = len(flagged)


def build_catalog(systeme: Path, iek: Path) -> dict:
    all_products = []
    source_summaries = []
    for vendor, archive in (("Systeme Electric", systeme), ("IEK", iek)):
        if not archive.exists():
            raise FileNotFoundError(f"Source archive not found: {archive}")
        files = source_files(archive)
        files["vendor"] = vendor
        if "sales" not in files or "moq" not in files:
            raise ValueError(f"Missing required monthly sales or MOQ workbook in {archive.name}")
        products: dict[str, dict] = {}
        extract_moq(files, products)
        extract_monthly_sales(files, products)
        extract_monthly_stock(files, products)
        if vendor == "Systeme Electric":
            extract_system_summary(files, products)
        else:
            extract_iek_inbound(files, products)
            # The IEK report is a complete list of open purchase documents.
            # A SKU absent from those lines has no recorded inbound quantity.
            if "iek_inbound" in files:
                for item in products.values():
                    item["inbound_known"] = True
        detect_document_outliers(files, products)
        catalog = sorted(products.values(), key=lambda item: (item["supplier"], item["sku"]))
        for item in catalog:
            item["monthly_sales"] = {month: round(value, 3) for month, value in sorted(item["monthly_sales"].items())}
            item["one_off_adjustments"] = dict(sorted(item["one_off_adjustments"].items()))
            if not item["name"]:
                item["name"] = item["supplier_sku"] or item["sku"]
        all_products.extend(catalog)
        source_summaries.append({
            "supplier": vendor,
            "archive": archive.name,
            "source_workbooks": {kind: entry[0] for kind, entry in files.items() if kind != "vendor"},
            "products": len(catalog),
            "with_sales_history": sum(bool(item["monthly_sales"]) for item in catalog),
            "with_stock": sum(item["stock_basis"] != "Остаток не найден в выгрузке" for item in catalog),
            "with_in_transit": sum(item["in_transit"] > 0 for item in catalog),
            "with_large_order_adjustment": sum(item["detected_large_orders"] > 0 for item in catalog),
        })
    return {
        "dataset_name": "Electrokomplekt supplier reports",
        "as_of_date": "2026-09-22",
        "data_mode": "supplier_report_aggregates",
        "sources": source_summaries,
        "data_limits": [
            "No anonymized client_id column was present. Large orders are grouped by invoice number and SKU, not by client.",
            "The supplied reports do not contain recorded stockout periods. Stockout CSV upload and what-if inputs are supported separately.",
            "The transaction ledgers do not include sale price. Recommendations are in item units, not currency.",
            "Supplier lead times are absent. The app uses a visible, editable planning assumption.",
            "IEK on-hand stock uses the September 2026 beginning-balance snapshot. Systeme Electric uses the free-stock figure from its 2026-09-22 summary.",
            "IEK has no explicit product-category field. Products without a matching MOQ row use a visible 1-unit fallback.",
        ],
        "products": all_products,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--systeme", type=Path, required=True, help="Path to Systeme electric.zip")
    parser.add_argument("--iek", type=Path, required=True, help="Path to IEK.zip")
    parser.add_argument("--output", type=Path, default=Path("data/demo_catalog.json"))
    args = parser.parse_args()
    catalog = build_catalog(args.systeme, args.iek)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(catalog, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({"output": str(args.output), "sources": catalog["sources"], "products": len(catalog["products"])}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
