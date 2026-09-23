"""Explicit adapters for the supplied IEK/Systeme ZIPs. No spreadsheet execution.
Raw files stay outside Git. Every consumed workbook is recorded in the manifest.
"""
from __future__ import annotations
import hashlib
import io
import json
import math
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import PurePosixPath
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook

MAX_UPLOAD = 25 * 1024 * 1024
MAX_EXPANDED = 512 * 1024 * 1024
MONTHS = {'янв': 1, 'фев': 2, 'мар': 3, 'апр': 4, 'май': 5, 'мая': 5, 'июн': 6, 'июл': 7, 'авг': 8, 'сен': 9, 'окт': 10, 'ноя': 11, 'дек': 12}


def num(x):
    if x is None or isinstance(x, bool):
        return None
    try:
        n = float(str(x).replace('\xa0', '').replace(' ', '').replace(',', '.'))
        return n if math.isfinite(n) else None
    except (ValueError, TypeError):
        return None


def text(x):
    return '' if x is None else str(x).strip()


def month(x):
    m = re.search(r'([а-яё]+)[.\s]+(20\d{2})', text(x).lower())
    if not m:
        return None
    n = MONTHS.get(m[1][:3])
    return f'{m[2]}-{n:02d}' if n else None


def digest(x):
    return hashlib.sha256(str(x).encode()).hexdigest()[:24]


def _archive_check(payload):
    if len(payload) > MAX_UPLOAD:
        raise ValueError('Файл превышает 25 МБ.')
    with ZipFile(io.BytesIO(payload)) as archive:
        if len(archive.infolist()) > 5000 or sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED:
            raise ValueError('Превышен допустимый распакованный объём архива.')
        for i in archive.infolist():
            path = PurePosixPath(i.filename.replace('\\', '/'))
            if path.is_absolute() or '..' in path.parts or i.flag_bits & 1:
                raise ValueError('Небезопасная структура или зашифрованный архив.')


def _kind(name):
    n = name.lower()
    if 'moq' in n:
        return 'moq'
    if 'динамика' in n:
        return 'transactions'
    if 'ежемесячные' in n and 'продажи' in n:
        return 'sales'
    if 'ежемесячные' in n and 'остатки' in n:
        return 'inventory'
    if 'сезонность' in n:
        return 'seasonality'
    if 'товар в пути' in n:
        return 'summary'
    if 'путь' in n:
        return 'inbound'
    return None


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = text(value)
    for fmt in ('%d.%m.%Y %H:%M:%S', '%d.%m.%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _headers(sheet):
    for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
        labels = [text(x).lower() for x in row]
        if any('код' in x for x in labels):
            return index, labels
        if index >= 10:
            break
    raise ValueError('Не найден заголовок с кодом товара.')


def _col(labels, *needles):
    return next((i for i, s in enumerate(labels) if all(n in s for n in needles)), None)


def _cell(row, col):
    return row[col] if col is not None and col < len(row) else None


def _sku(raw):
    if raw is None:
        return ''
    if not isinstance(raw, str):
        raise ValueError('Код товара хранится как число. Нужна текстовая выгрузка, сохраняющая ведущие нули.')
    value = raw.strip()
    return '' if value.lower().startswith(('итого', 'всего', '#')) else value


def _new(sku, supplier):
    return {'sku': sku, 'supplier': supplier, 'supplier_sku': '', 'warehouse': 'Сводный склад (выгрузка)',
            'name': '', 'unit': '', 'category': '', 'category_known': False, 'current_stock': None,
            'stock_known': False, 'stock_snapshot': None, 'inbound_known': False, 'in_transit': None,
            'inbound_shipments': [], 'monthly_sales': {}, 'monthly_stock': {}, 'transactions': [],
            'moq_known': False, 'order_multiple': 1, 'min_order': 0, 'approval_blockers': [], 'data_warnings': [], 'provenance': {}}


def parse_supplier_archive(payload: bytes, supplier: str, as_of: date) -> dict:
    if supplier not in ('IEK', 'Systeme Electric'):
        raise ValueError('Неизвестный поставщик.')
    _archive_check(payload)
    files = {}
    with ZipFile(io.BytesIO(payload)) as archive:
        for member in archive.infolist():
            if member.filename.lower().endswith('.xlsx'):
                kind = _kind(member.filename)
                if not kind:
                    raise ValueError('Неизвестная Excel-выгрузка: ' + PurePosixPath(member.filename).name)
                if kind in files:
                    raise ValueError('Повторная выгрузка одного типа: ' + kind)
                files[kind] = (PurePosixPath(member.filename).name, archive.read(member))
    required = {'moq', 'transactions', 'sales', 'inventory', 'seasonality', 'inbound' if supplier == 'IEK' else 'summary'}
    if not required.issubset(files):
        raise ValueError('Не хватает типов выгрузок: ' + ', '.join(sorted(required - files.keys())))
    products, manifests, warnings, vendor_factors = {}, [], [], {}
    # A newer summary must take precedence over historical beginning balances.
    order = ['moq', 'sales', 'inventory', 'summary', 'inbound', 'seasonality', 'transactions']
    for kind in order:
        if kind not in files:
            continue
        filename, data = files[kind]
        _archive_check(data)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheet = workbook['TDSheet'] if kind == 'summary' and 'TDSheet' in workbook.sheetnames else workbook.worksheets[0]
            if kind == 'seasonality':
                years = []
                for row in sheet.iter_rows(values_only=True):
                    values = list(row)
                    for i, x in enumerate(values):
                        year = num(x)
                        if year is None or year != int(year) or not 2000 <= year < as_of.year:
                            continue
                        vals = [num(v) for v in values[i+1:i+13]]
                        if len(vals) == 12 and all(v is not None and v >= 0 for v in vals) and sum(vals) > 0:
                            avg = sum(vals) / 12
                            years.append([v / avg for v in vals])
                            break
                if years:
                    vendor_factors = {str(m+1): sum(y[m] for y in years) / len(years) for m in range(12)}
                else:
                    warnings.append('Сезонность: не найдены полные фактические годы; прогнозные строки не использованы.')
                manifests.append({'kind': kind, 'file': filename, 'sheet': sheet.title, 'sha256': hashlib.sha256(data).hexdigest(), 'full_years': len(years)})
                continue
            header, labels = _headers(sheet)
            if any(any(term in label for term in ('фио', 'телефон', 'email', 'e-mail', 'контрагент.наименование', 'имя клиента')) for label in labels):
                raise ValueError('Обнаружены возможные персональные данные; требуется обезличенная выгрузка.')
            code_col = _col(labels, 'код')
            name_col = next((i for i, x in enumerate(labels) if ('наименован' in x or 'номенклатур' in x) and i != code_col), None)
            months = [(i, month(s)) for i, s in enumerate(labels) if month(s)]
            seen_skus, seen_events = set(), set()
            records = 0
            for row_index, row in enumerate(sheet.iter_rows(min_row=header+1, values_only=True), header+1):
                sku = _sku(_cell(row, code_col))
                if not sku:
                    continue
                p = products.setdefault(sku, _new(sku, supplier))
                if kind != 'transactions' and sku in seen_skus:
                    p['approval_blockers'].append('Повторный код в ' + filename + '; требуется сверка, строки не суммировались.')
                    continue
                seen_skus.add(sku)
                records += 1
                p['provenance'][kind] = {'file': filename, 'sheet': sheet.title, 'row': row_index}
                name = text(_cell(row, name_col))
                if name:
                    p['name'] = name
                article = _cell(row, _col(labels, 'артикул'))
                if article is not None:
                    article = text(article)
                    if p['supplier_sku'] and p['supplier_sku'] != article:
                        p['approval_blockers'].append('Конфликт артикула поставщика между выгрузками.')
                    p['supplier_sku'] = article
                unit_col = next((i for i, s in enumerate(labels) if s in ('ед.', 'ед.изм', 'ед. изм.')), None)
                if _cell(row, unit_col):
                    p['unit'] = text(_cell(row, unit_col))
                if kind == 'moq':
                    col = next((i for i, s in enumerate(labels) if 'кратност' in s or 'разр' in s), None)
                    value = num(_cell(row, col))
                    if value is not None and value > 0:
                        if 'кратност' in labels[col]:
                            p['order_multiple'] = value
                            p['moq_known'] = True
                        else:
                            p['min_order'] = value
                            p['approval_blockers'].append('Подтвердить смысл «Мин. разр. к отгр.»: минимум или кратность.')
                elif kind == 'sales':
                    for i, period in months:
                        value = num(_cell(row, i))
                        p['monthly_sales'][period] = value
                        if value is None:
                            p['data_warnings'].append('Пустые продажи оставлены неизвестными, не нулями.')
                        if value is not None and value < 0:
                            p['approval_blockers'].append('Есть отрицательные месячные продажи; нужна сверка возвратов.')
                elif kind == 'inventory':
                    p['monthly_stock'] = {period: num(_cell(row, i)) for i, period in months}
                    usable = [(period, num(_cell(row, i))) for i, period in months if period <= as_of.strftime('%Y-%m')]
                    if usable:
                        period, value = max(usable)
                        p.update(current_stock=value, stock_known=value is not None and value >= 0,
                                 stock_snapshot=period + '-01', stock_basis='Начальный месячный остаток; не текущий срез')
                elif kind == 'summary':
                    free_col = _col(labels, 'свободный остаток')
                    stock_col = next((i for i, s in enumerate(labels) if s == 'остаток'), None)
                    raw_free = num(_cell(row, free_col))
                    total, reserved = num(_cell(row, stock_col)), num(_cell(row, _col(labels, 'зарезервировано')))
                    p['reserved'] = reserved
                    p['reported_stock'] = total
                    p['current_stock'] = raw_free
                    p['stock_known'] = raw_free is not None and raw_free >= 0
                    p['stock_snapshot'] = as_of.isoformat()
                    p['stock_basis'] = 'Свободный остаток; резерв повторно не вычитается'
                    category = text(_cell(row, _col(labels, 'категория')))
                    p.update(category=category, category_known=bool(category))
                    growth = num(_cell(row, _col(labels, 'роста')))
                    p['growth_source_value'] = growth
                    if growth is not None and 1 + growth > 0:
                        p['growth_factor'] = 1 + growth
                        p['approval_blockers'].append('Подтвердить: «Кэф. Роста» интерпретирован как прирост, 1 + значение, без повторного тренда.')
                    p['legacy_seasonality'] = num(_cell(row, _col(labels, 'сез-ти')))
                    for i, period in months:
                        val = num(_cell(row, i))
                        prior = p['monthly_sales'].get(period)
                        if prior is not None and val is not None and abs(prior-val) > 1e-6:
                            p['approval_blockers'].append('Месячные продажи отличаются от сводного отчёта; нужна сверка.')
                if kind in ('summary', 'inbound'):
                    cols = [i for i, s in enumerate(labels) if 'поступление' in s or 'в пути' in s]
                    complete = bool(cols)
                    for i in cols:
                        quantity = num(_cell(row, i))
                        if quantity is None:
                            complete = False
                            continue
                        if quantity < 0:
                            p['approval_blockers'].append('Отрицательное количество поставки.')
                            continue
                        dates = re.findall(r'(\d{2})\.(\d{2})\.(20\d{2})', labels[i])
                        short = re.findall(r'(\d{2})\.(\d{2})', labels[i])
                        eta = None
                        try:
                            if dates:
                                d, m, y = dates[-1]
                                eta = date(int(y), int(m), int(d)).isoformat()
                            elif short:
                                d, m = short[-1]
                                eta = date(as_of.year, int(m), int(d)).isoformat()
                        except ValueError:
                            pass
                        if quantity:
                            p['inbound_shipments'].append({'quantity': quantity, 'eta': eta, 'status': 'confirmed', 'source_column': labels[i]})
                    p['inbound_known'] = complete
                    p['in_transit'] = sum(s['quantity'] for s in p['inbound_shipments']) if complete else None
                    if not complete:
                        p['data_warnings'].append('Пустые ячейки поставок требуют подтверждения, что они означают ноль.')
                elif kind == 'transactions':
                    qty = num(_cell(row, _col(labels, 'количество')))
                    d = _date(_cell(row, next((i for i, s in enumerate(labels) if s == 'дата'), None)))
                    description = text(_cell(row, _col(labels, 'документ'))).lower()
                    if qty is None or d is None or d > as_of or qty >= 0:
                        continue
                    if not any(s in description for s in ('расходная накладная', 'реализация', 'продажа')):
                        continue
                    document = text(_cell(row, _col(labels, 'номер')))
                    warehouse = text(_cell(row, _col(labels, 'склад')))
                    event_id = digest((supplier, sku, document, d.isoformat(), qty, warehouse))
                    if event_id in seen_events:
                        p['data_warnings'].append('Повторная операция обнаружена; нужна сверка идентификатора строки документа.')
                        continue
                    seen_events.add(event_id)
                    customer_col = next((i for i, s in enumerate(labels) if s in ('customer_id', 'client_id', 'обезличенный id клиента')), None)
                    client = text(_cell(row, customer_col))
                    if client and not re.fullmatch(r'[A-Za-z0-9_-]{1,80}', client):
                        raise ValueError('ID клиента должен быть обезличенным кодом, а не именем или контактом.')
                    p['transactions'].append({'event_id': event_id, 'document_id': digest((supplier, document)), 'customer_id': digest((supplier, client)) if client else None,
                                              'date': d.isoformat(), 'quantity': -qty, 'warehouse': warehouse})
            manifests.append({'kind': kind, 'file': filename, 'sheet': sheet.title, 'sha256': hashlib.sha256(data).hexdigest(), 'rows': records})
        finally:
            workbook.close()
    for p in products.values():
        p['name'] = p['name'] or p['sku']
        p['seasonality_factors'] = vendor_factors
        if vendor_factors:
            p['data_warnings'].append('Сезонный профиль поставщика: агрегированные показатели, не отдельный SKU; нужна проверка на backtest.')
        if 'бухт' in p['name'].lower() and ('метр' in p['name'].lower() or p['unit'] in ('м', 'м.')):
            p['purchase_unit'] = 'бухта'
            p['conversion_confirmed'] = False
        else:
            p['purchase_unit'] = p['unit']
            p['units_per_purchase'] = 1
        if not p['unit']:
            p['approval_blockers'].append('Единица измерения не установлена.')
        if any(v is None for v in p['monthly_sales'].values()):
            p['approval_blockers'].append('Подтвердить смысл пустых месячных продаж; отсутствующие периоды не считались нулями.')
        p['data_warnings'] = sorted(set(p['data_warnings']))
        p['approval_blockers'] = sorted(set(p['approval_blockers']))
    return {'schema_version': 2, 'as_of_date': as_of.isoformat(), 'data_mode': 'supplier_reports',
            'products': list(products.values()), 'sources': [{'supplier': supplier, 'products': len(products),
            'with_sales_history': sum(bool(p['monthly_sales']) for p in products.values()), 'with_stock': sum(p['stock_known'] for p in products.values()), 'workbooks': manifests}],
            'category_policies': {}, 'data_limits': warnings + ['Точные интервалы stockout и client_id в исходном пакете не гарантированы.',
            'Неизвестные поля не заменяются нулём. Запуск расчёта не является разрешением отправить заказ.',
            'Нет подтверждённой спецификации импорта 1С: экспорт является черновым обменным форматом.']}


def merge_catalogs(catalogs):
    if not catalogs or len({c['as_of_date'] for c in catalogs}) != 1:
        raise ValueError('Нужны каталоги на одну явно заданную дату.')
    products = [p for c in catalogs for p in c['products']]
    keys = [(p['supplier'], p['sku'], p.get('warehouse', '')) for p in products]
    if len(keys) != len(set(keys)):
        raise ValueError('Повторный импорт поставщика/артикула/склада. Набор отклонён, чтобы не удваивать продажи.')
    return {**catalogs[0], 'products': products, 'sources': [s for c in catalogs for s in c['sources']], 'data_limits': sorted({x for c in catalogs for x in c['data_limits']})}
